"""Exportaciones en memoria para la aplicación web."""

from io import BytesIO
from html import escape
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BRAND_BLUE = colors.HexColor('#123B6D')
BRAND_NAVY = colors.HexColor('#08111F')
LIGHT_GRAY = colors.HexColor('#E9EDF2')


def business_as_dict(profile) -> dict:
    fields = ('business_name', 'owner_name', 'phone', 'whatsapp', 'email', 'address',
              'tax_id', 'website')
    return {field: getattr(profile, field, '') or '' for field in fields}


def _business_lines(business: dict) -> list[str]:
    contact = ' · '.join(value for value in (
        business.get('phone'), business.get('whatsapp'), business.get('email')) if value)
    return [value for value in (
        business.get('business_name') or 'Mi negocio', business.get('owner_name'),
        business.get('address'), contact,
        f"RFC: {business['tax_id']}" if business.get('tax_id') else '',
        business.get('website'),
    ) if value]


def generate_pdf_bytes(business: dict, title: str, subtitle: str,
                       sections: Iterable[tuple[str, list[dict]]]) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=16 * mm,
                                 rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    centered = ParagraphStyle('Centered', parent=styles['Normal'], alignment=TA_CENTER,
                              textColor=BRAND_NAVY, leading=14)
    story = []
    business_lines = _business_lines(business)
    story.append(Paragraph(f'<b>{business_lines[0]}</b>', styles['Title']))
    if len(business_lines) > 1:
        story.append(Paragraph('<br/>'.join(business_lines[1:]), centered))
    story.extend((Spacer(1, 8), Paragraph(f'<b>{title}</b>', styles['Heading1']),
                  Paragraph(subtitle, styles['Normal']), Spacer(1, 10)))

    for section_title, rows in sections:
        story.append(Paragraph(f'<b>{section_title}</b>', styles['Heading2']))
        if not rows:
            story.append(Paragraph('Sin registros.', styles['Normal']))
            story.append(Spacer(1, 8))
            continue
        headers = list(rows[0].keys())
        cell_style = ParagraphStyle('TableCell', parent=styles['Normal'], fontSize=7,
                                    leading=9, wordWrap='CJK')
        values = [[Paragraph(f'<b>{escape(str(header))}</b>', cell_style)
                   for header in headers]]
        values.extend([
            [Paragraph(escape(str(row.get(header, ''))), cell_style) for header in headers]
            for row in rows
        ])
        available_width = A4[0] - 32 * mm
        table = Table(values, repeatRows=1, hAlign='LEFT',
                      colWidths=[available_width / len(headers)] * len(headers))
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), BRAND_BLUE),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#B8C1CC')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), (colors.white, LIGHT_GRAY)),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.extend((table, Spacer(1, 10)))
    document.build(story)
    return output.getvalue()


def generate_excel_bytes(business: dict, title: str, subtitle: str,
                         sections: Iterable[tuple[str, list[dict]]]) -> bytes:
    workbook = Workbook()
    information = workbook.active
    information.title = 'Información'
    information.append([business.get('business_name') or 'Mi negocio'])
    information.append([title])
    information.append([subtitle])
    for line in _business_lines(business)[1:]:
        information.append([line])
    information['A1'].font = Font(bold=True, size=16, color='123B6D')
    information['A2'].font = Font(bold=True, size=13)
    information.column_dimensions['A'].width = 70

    used_names = {'Información'}
    for section_title, rows in sections:
        base_name = section_title[:31] or 'Datos'
        sheet_name = base_name
        suffix = 2
        while sheet_name in used_names:
            sheet_name = f'{base_name[:27]} {suffix}'
            suffix += 1
        used_names.add(sheet_name)
        sheet = workbook.create_sheet(sheet_name)
        if not rows:
            sheet.append(['Sin registros'])
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='123B6D')
            cell.alignment = Alignment(horizontal='center')
        for row in rows:
            sheet.append([row.get(header, '') for header in headers])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(45, max(len(str(cell.value or '')) for cell in column_cells) + 2)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
